from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from pypdf import PdfReader
from sqlalchemy.orm import Session
from typer.testing import CliRunner

from app.cli.commands import app
from app.documents.engine import DocumentBuilder, _deliverable_content, _deliverable_name
from app.errors import AppError
from app.llm.gemini import generation_schema
from app.llm.providers import MockProvider
from app.llm.schema import Solution
from app.llm.service import solve_context
from app.persistence.models import GeneratedArtifact, SolutionRecord
from tests.test_phase4 import context_for  # noqa: F401
from tests.test_phase4 import phase4 as phase4

FORMAT_FIXTURE = """# Infraestrutura de TI

A **infraestrutura de TI** inclui *redes* e o termo `servidor Linux`.

## Principais componentes

- Servidores
- Redes
- Armazenamento

1. Planejamento
2. Implantação

```python
if 2 < 3:
    print("Olá, mundo")
```

| Componente | Função |
| --- | --- |
| Rede | Comunicação |
"""


def prepared_builder(phase4):
    settings, engine, _ = phase4
    solution = solve_context(engine, context_for(phase4), MockProvider())
    with Session(engine) as session, session.begin():
        row = session.get(SolutionRecord, solution.id)
        assert row is not None
        row.response = {
            "assignment_types": ["LONG_FORM"],
            "understanding": "",
            "answer": FORMAT_FIXTURE,
            "question_answers": [],
            "artifacts": [],
        }
    configured = settings.model_copy(
        update={"generated_root": settings.database_file.parent / "format-output"}
    )
    return configured, engine, DocumentBuilder(engine, configured)


def test_requested_html_is_materialized_alongside_pdf(phase4):
    settings, engine, _ = phase4
    solution = solve_context(engine, context_for(phase4), MockProvider())
    with Session(engine) as session, session.begin():
        row = session.get(SolutionRecord, solution.id)
        assert row and row.response
        response = dict(row.response)
        response["assignment_types"] = ["PROGRAMMING"]
        response["artifacts"] = [
            {
                "kind": "CODE",
                "title": "index.html",
                "specification": "<!doctype html>\n<html><body>OK</body></html>",
            }
        ]
        row.response = response

    artifacts = DocumentBuilder(
        engine,
        settings.model_copy(update={"generated_root": settings.database_file.parent / "out"}),
    ).build_all(1)

    assert [artifact.artifact_type for artifact in artifacts] == ["PDF", "HTML"]
    html = Path(artifacts[1].local_path)
    assert html.name == "index.html"
    assert html.read_text(encoding="utf-8").startswith("<!doctype html>")
    assert all(artifact.status == "UPLOAD_PENDING" for artifact in artifacts)


def test_deliverable_rejects_paths_executables_and_nul():
    assert _deliverable_name("index.html") == "index.html"
    assert _deliverable_name("../index.html") is None
    assert _deliverable_name("run.exe") is None
    assert _deliverable_name("token.json") is None
    assert _deliverable_content("```html\n<p>OK</p>\n```") == "<p>OK</p>\n"
    assert _deliverable_content("bad\x00data") is None


def test_pdf_format_renders_markdown_structure_without_tokens(phase4):
    _, _, builder = prepared_builder(phase4)
    artifact = builder.build(1, override="academic-report", output_format="pdf")

    extracted = "\n".join(
        page.extract_text() or "" for page in PdfReader(artifact.local_path).pages
    )
    assert "Infraestrutura de TI" in extracted
    assert "Principais componentes" in extracted
    assert "infraestrutura de TI" in extracted
    assert "servidor Linux" in extracted and 'print("Olá, mundo")' in extracted
    assert "2 < 3" in extracted
    assert "Servidores" in extracted and "Planejamento" in extracted
    assert "Componente" in extracted and "Comunicação" in extracted
    assert "**" not in extracted and "```" not in extracted
    assert not any(line.lstrip().startswith("#") for line in extracted.splitlines())


def test_docx_is_valid_and_preserves_heading_emphasis_lists_and_code(phase4):
    _, _, builder = prepared_builder(phase4)
    artifact = builder.build(1, override="academic-report", output_format="docx")

    document = Document(artifact.local_path)
    paragraphs = document.paragraphs
    assert any(p.text == "Infraestrutura de TI" and p.style.name == "Heading 1" for p in paragraphs)
    assert any(
        p.text == "Principais componentes" and p.style.name == "Heading 2" for p in paragraphs
    )
    assert any(
        run.text == "infraestrutura de TI" and run.bold for p in paragraphs for run in p.runs
    )
    assert any(run.text == "redes" and run.italic for p in paragraphs for run in p.runs)
    assert any(p.text == "Servidores" and p.style.name == "List Bullet" for p in paragraphs)
    assert any(p.text == "1. Planejamento" or p.text == "Planejamento" for p in paragraphs)
    assert any(p.style.name == "List Number" for p in paragraphs)
    assert any('print("Olá, mundo")' in p.text and "2 < 3" in p.text for p in paragraphs)
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "Componente"
    assert document.tables[0].cell(1, 1).text == "Comunicação"
    assert any(
        run.text == "servidor Linux" and run.font.name == "Consolas"
        for paragraph in paragraphs
        for run in paragraph.runs
    )
    assert all("**" not in p.text and "```" not in p.text for p in paragraphs)


def test_txt_is_clean_and_markdown_tokens_are_removed(phase4):
    _, _, builder = prepared_builder(phase4)
    artifact = builder.build(1, override="academic-report", output_format="txt")
    text = Path(artifact.local_path).read_text(encoding="utf-8")

    assert "Infraestrutura de TI" in text and "Principais componentes" in text
    assert "infraestrutura de TI" in text and "redes" in text and "servidor Linux" in text
    assert "Servidores" in text and "1. Planejamento" in text
    assert 'print("Olá, mundo")' in text and "2 < 3" in text and "Comunicação" in text
    assert "**" not in text and "```" not in text and "`" not in text
    assert not any(line.lstrip().startswith("#") for line in text.splitlines())


def test_md_preserves_markdown_syntax(phase4):
    _, _, builder = prepared_builder(phase4)
    artifact = builder.build(1, override="academic-report", output_format="md")
    text = Path(artifact.local_path).read_text(encoding="utf-8")

    assert "# Infraestrutura de TI" in text
    assert "**infraestrutura de TI**" in text
    assert "*redes*" in text and "`servidor Linux`" in text
    assert "## Principais componentes" in text
    assert "- Servidores" in text and "1. Planejamento" in text
    assert '```python\nif 2 < 3:\n    print("Olá, mundo")\n```' in text
    assert "| Componente | Função |" in text


def test_code_assignment_removes_outer_fence_without_changing_code(phase4):
    _, engine, builder = prepared_builder(phase4)
    with Session(engine) as session, session.begin():
        row = session.query(SolutionRecord).one()
        row.response = {
            "assignment_types": ["PROGRAMMING"],
            "understanding": "",
            "answer": "Código de exemplo.",
            "deliverable": "",
            "question_answers": [],
            "artifacts": [
                {
                    "kind": "CODE",
                    "title": "main.py",
                    "specification": "```python\ndef main():\n    print('ok')\n```",
                }
            ],
        }
    artifact = builder.build(1, override="code-assignment", output_format="txt")
    text = Path(artifact.local_path).read_text(encoding="utf-8")

    assert "def main():" in text and "    print('ok')" in text
    assert "```" not in text
    assert "Entendimento" not in text and "Resposta" not in text


def test_academic_deliverable_excludes_internal_solution_fields_and_follows_task(phase4, caplog):
    _, engine, builder = prepared_builder(phase4)
    deliverable = """# Estudo de caso: infraestrutura para clínica regional

Uma clínica com três unidades precisa manter prontuários disponíveis e proteger dados
sensíveis. A proposta combina serviços de identidade, prontuário eletrônico, DNS,
monitoramento e cópias de segurança.

## Sistemas, usuários e equipamentos

Recepcionistas consultam agendas; profissionais de saúde registram atendimentos; a equipe
de TI administra contas e auditoria. Computadores gerenciados acessam os sistemas por uma
rede segmentada, com Wi-Fi separado para visitantes.

## Hospedagem, continuidade e virtualização

O prontuário pode operar em nuvem privada gerenciada, enquanto integrações locais permanecem
no datacenter da unidade principal. Máquinas virtuais isolam banco de dados, aplicação e
monitoramento. Cópias criptografadas em região distinta e um procedimento de contingência
permitem restaurar o serviço após falha.
"""
    with Session(engine) as session, session.begin():
        row = session.query(SolutionRecord).one()
        row.response = {
            "assignment_types": ["RESEARCH"],
            "understanding": "O enunciado solicita discutir infraestrutura. PLANEJAMENTO INTERNO.",
            "answer": "PRÉVIA INTERNA: proposta para uma clínica.",
            "deliverable": deliverable,
            "question_answers": [],
            "artifacts": [],
        }
    with caplog.at_level("DEBUG", logger="app.documents.engine"):
        artifact = builder.build(1, override="academic-report", output_format="docx")
    document = Document(artifact.local_path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    document_text = "\n".join(paragraphs)

    assert any(
        paragraph.text == "Estudo de caso: infraestrutura para clínica regional"
        and paragraph.style.name == "Heading 1"
        for paragraph in document.paragraphs
    )
    assert "Sistemas, usuários e equipamentos" in document_text
    assert "Hospedagem, continuidade e virtualização" in document_text
    assert "Entendimento" not in document_text and "Resposta" not in document_text
    assert "O enunciado solicita" not in document_text
    assert "PRÉVIA INTERNA" not in document_text and "PLANEJAMENTO INTERNO" not in document_text
    record = next(record for record in caplog.records if "deliverable_strategy" in record.message)
    assert "artifact=academic_report" in record.message
    assert "solution_version=1" in record.message
    assert "deliverable_strategy=structured" in record.message
    assert "clínica" not in record.message and "Estudo de caso" not in record.message


def test_academic_deliverable_keeps_explicit_analysis_heading(phase4):
    _, engine, builder = prepared_builder(phase4)
    with Session(engine) as session, session.begin():
        row = session.query(SolutionRecord).one()
        row.response = {
            "assignment_types": ["LONG_FORM"],
            "understanding": "Planejamento interno.",
            "answer": "Prévia.",
            "deliverable": "## Análise\n\nA atividade solicita uma seção com este título.",
            "question_answers": [],
            "artifacts": [],
        }
    artifact = builder.build(1, override="academic-report", output_format="txt")
    text = Path(artifact.local_path).read_text(encoding="utf-8")

    assert "Análise" in text and "A atividade solicita uma seção" in text
    assert "Planejamento interno" not in text and "Prévia." not in text


def test_academic_deliverable_can_number_ten_requested_items(phase4):
    _, engine, builder = prepared_builder(phase4)
    lines = "\n".join(f"{index}. Aspecto {index}" for index in range(1, 11))
    with Session(engine) as session, session.begin():
        row = session.query(SolutionRecord).one()
        row.response = {
            "assignment_types": ["LONG_FORM"],
            "understanding": "Liste dez itens.",
            "answer": "Prévia.",
            "deliverable": lines,
            "question_answers": [],
            "artifacts": [],
        }
    artifact = builder.build(1, override="academic-report", output_format="txt")
    text = Path(artifact.local_path).read_text(encoding="utf-8")

    assert all(f"{index}. Aspecto {index}" in text for index in range(1, 11))
    assert "Entendimento" not in text and "Liste dez itens" not in text


def test_new_code_template_excludes_internal_understanding_and_keeps_source(phase4):
    _, engine, builder = prepared_builder(phase4)
    with Session(engine) as session, session.begin():
        row = session.query(SolutionRecord).one()
        row.response = {
            "assignment_types": ["PROGRAMMING"],
            "understanding": "Planejar a implementação.",
            "answer": "A implementação usa uma função principal.",
            "deliverable": "",
            "question_answers": [],
            "artifacts": [
                {
                    "kind": "CODE",
                    "title": "main.py",
                    "specification": "def main():\n    return 42\n",
                }
            ],
        }
    artifact = builder.build(1, override="code-assignment", output_format="txt")
    text = Path(artifact.local_path).read_text(encoding="utf-8")

    assert "def main():" in text and "    return 42" in text
    assert "Planejar a implementação" not in text
    assert "Resposta" not in text


def test_legacy_solution_without_deliverable_still_generates_unchanged(phase4):
    _, engine, builder = prepared_builder(phase4)
    with Session(engine) as session, session.begin():
        row = session.query(SolutionRecord).one()
        row.response = {
            "assignment_types": ["LONG_FORM"],
            "understanding": "Interpretação antiga preservada.",
            "answer": "Resposta antiga preservada.",
            "question_answers": [],
            "artifacts": [],
        }
    artifact = builder.build(1, override="academic-report", output_format="txt")
    text = Path(artifact.local_path).read_text(encoding="utf-8")

    assert "Entendimento" in text and "Interpretação antiga preservada" in text
    assert "Resposta" in text and "Resposta antiga preservada" in text


def test_new_academic_markdown_renders_in_pdf_and_docx_without_internal_text(phase4):
    _, engine, builder = prepared_builder(phase4)
    with Session(engine) as session, session.begin():
        row = session.query(SolutionRecord).one()
        row.response = {
            "assignment_types": ["RESEARCH"],
            "understanding": "INTERNO: O enunciado solicita um relatório.",
            "answer": "PRÉVIA INTERNA.",
            "deliverable": FORMAT_FIXTURE,
            "question_answers": [],
            "artifacts": [],
        }
    pdf = builder.build(1, override="academic-report", output_format="pdf")
    docx = builder.build(1, override="academic-report", output_format="docx")
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf.local_path).pages)
    docx_text = "\n".join(p.text for p in Document(docx.local_path).paragraphs)

    for rendered in (pdf_text, docx_text):
        assert "Principais componentes" in rendered
        assert "servidor Linux" in rendered and "Olá, mundo" in rendered
        assert "INTERNO" not in rendered and "PRÉVIA INTERNA" not in rendered
        assert "Entendimento" not in rendered and "Resposta" not in rendered
        assert "**" not in rendered and "```" not in rendered


def test_question_answer_template_uses_question_pairs_not_deliverable():
    from app.documents.engine import Template, _document_sections

    sections = _document_sections(
        Template.QUESTION_ANSWER,
        {
            "understanding": "Não incluir.",
            "answer": "Fallback.",
            "deliverable": "",
            "question_answers": [
                {"question_id": "A", "answer": "Resposta A."},
                {"question_id": "B", "answer": "Resposta B."},
            ],
        },
    )

    assert sections == [("1. A", "Resposta A."), ("2. B", "Resposta B.")]


def test_template_specific_prompts_define_deliverable_contract(phase4):
    from app.llm.prompt import build_system_prompt

    context = context_for(phase4)
    academic = build_system_prompt(context.model_copy(update={"assignment_types": ["RESEARCH"]}))
    questions = build_system_prompt(
        context.model_copy(update={"assignment_types": ["SHORT_ANSWER"]})
    )
    code = build_system_prompt(context.model_copy(update={"assignment_types": ["PROGRAMMING"]}))

    assert "FORMATO DE TRABALHO ACADÊMICO" in academic
    assert "Não repita o enunciado" in academic and "deliverable" in academic
    assert "FORMATO DE PERGUNTAS E RESPOSTAS" in questions
    assert 'use deliverable=""' in questions
    assert "FORMATO DE PROGRAMAÇÃO" in code
    assert "Não altere o código" in code and 'deliverable=""' in code


def test_deliverable_schema_is_required_simple_string_for_providers():
    schema = generation_schema(Solution.model_json_schema())

    assert schema["properties"]["deliverable"]["type"] == "string"
    assert "deliverable" in schema["required"]


@pytest.mark.parametrize("invalid", ["xyz", "../../arquivo", "/tmp/test.pdf", "foo/bar.docx"])
def test_format_rejects_unknown_values_and_paths(phase4, invalid):
    _, _, builder = prepared_builder(phase4)
    with pytest.raises(AppError, match="pdf, docx, txt ou md"):
        builder.build(1, output_format=invalid)


def test_cli_template_and_format_are_independent_and_legacy_defaults_to_pdf(phase4):
    settings, engine, _ = phase4
    solution = solve_context(engine, context_for(phase4), MockProvider())
    with Session(engine) as session, session.begin():
        row = session.get(SolutionRecord, solution.id)
        assert row is not None
        row.response = {
            "assignment_types": ["LONG_FORM"],
            "understanding": "",
            "answer": FORMAT_FIXTURE,
            "question_answers": [],
            "artifacts": [],
        }
    settings = settings.model_copy(
        update={"generated_root": settings.database_file.parent / "format-cli"}
    )
    with patch("app.cli.commands.load_settings", return_value=settings):
        result = CliRunner().invoke(
            app,
            ["generate", "1", "--template", "academic-report", "--format", "docx"],
        )
    assert result.exit_code == 0, result.output
    assert "DOCX:" in result.output
    with Session(engine) as session:
        artifact = session.query(GeneratedArtifact).filter_by(artifact_type="DOCX").one()
        assert artifact.template == "ACADEMIC_REPORT"
        assert Document(artifact.local_path).paragraphs[0].style.name == "Title"


def test_generate_help_and_invalid_format_are_clear(phase4):
    settings, _, _ = phase4
    runner = CliRunner()
    help_result = runner.invoke(app, ["generate", "--help"])
    assert help_result.exit_code == 0
    flattened_help = " ".join(help_result.output.split())
    assert "--format" in flattened_help
    assert all(value in flattened_help for value in ("pdf", "docx", "txt", "md"))
    with patch("app.cli.commands.load_settings", return_value=settings):
        result = runner.invoke(app, ["generate", "1", "--format", "../../escape"])
    assert result.exit_code == 1
    assert "Formato inválido" in result.output
    assert "pdf, docx, txt ou md" in result.output
