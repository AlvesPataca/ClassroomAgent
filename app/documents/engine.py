from __future__ import annotations

import hashlib
import re
from datetime import UTC, datetime
from enum import StrEnum
from pathlib import Path
from typing import Any

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import func, select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session

from app.config import Settings
from app.documents.formatting import Content, Node, clean_text, parse_markdown, text_content
from app.errors import AppError
from app.persistence.models import AssignmentRecord, GeneratedArtifact, SolutionRecord


class Template(StrEnum):
    QUESTION_ANSWER = "question-answer"
    ACADEMIC_REPORT = "academic-report"
    CODE_ASSIGNMENT = "code-assignment"


class OutputFormat(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    TXT = "txt"
    MD = "md"


FORMAT_EXTENSIONS = {
    OutputFormat.PDF: ".pdf",
    OutputFormat.DOCX: ".docx",
    OutputFormat.TXT: ".txt",
    OutputFormat.MD: ".md",
}

FORMAT_ARTIFACT_TYPES = {
    OutputFormat.PDF: "PDF",
    OutputFormat.DOCX: "DOCX",
    OutputFormat.TXT: "TEXT",
    OutputFormat.MD: "MARKDOWN",
}


TEXT_DELIVERABLES = {
    ".html": "HTML",
    ".htm": "HTML",
    ".css": "CSS",
    ".js": "JAVASCRIPT",
    ".json": "JSON",
    ".xml": "XML",
    ".py": "PYTHON",
    ".java": "JAVA",
    ".c": "C",
    ".h": "HEADER",
    ".cpp": "CPP",
    ".sql": "SQL",
    ".md": "MARKDOWN",
    ".txt": "TEXT",
    ".csv": "CSV",
}


def _safe(value: str, fallback: str = "document") -> str:
    value = re.sub(r"[<>:\\|?*\x00-\x1f]", " ", value or "")
    value = re.sub(r"[\s./]+", " ", value).strip(" .")
    if not value or value.upper() in {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        *(f"COM{i}" for i in range(1, 10)),
    }:
        return fallback
    return value[:120]


def _esc(value: Any) -> str:
    from xml.sax.saxutils import escape

    return escape(str(value or "")).replace("\n", "<br/>")


def _deliverable_name(value: object) -> str | None:
    """Accept one harmless filename, never a path supplied by the model."""
    if not isinstance(value, str) or Path(value).name != value:
        return None
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", value).strip(" .")[:120]
    if not name or name.lower() in {".env", "token.json", "credentials.json", "classroom.db"}:
        return None
    return name if Path(name).suffix.lower() in TEXT_DELIVERABLES else None


def _deliverable_content(value: object) -> str | None:
    if not isinstance(value, str) or not value.strip() or "\x00" in value:
        return None
    text = value.strip()
    fenced = re.fullmatch(r"```[A-Za-z0-9_+.-]*\s*\n(.*)\n```", text, re.DOTALL)
    if fenced:
        text = fenced.group(1)
    return text + ("" if text.endswith("\n") else "\n")


def _inline_reportlab(content: list[Content]) -> str:
    from xml.sax.saxutils import escape

    output: list[str] = []
    for item in content:
        if isinstance(item, str):
            output.append(escape(item))
            continue
        inner = _inline_reportlab(item.children)
        if item.tag in {"strong", "b"}:
            inner = f"<b>{inner}</b>"
        elif item.tag in {"em", "i"}:
            inner = f"<i>{inner}</i>"
        elif item.tag == "code":
            inner = f"<font name='Courier'>{inner}</font>"
        elif item.tag in {"del", "s", "strike"}:
            inner = f"<strike>{inner}</strike>"
        elif item.tag == "br":
            inner += "<br/>"
        output.append(inner)
    return "".join(output)


def _pdf_blocks(nodes: list[Node], styles: Any, *, depth: int = 0) -> list[Any]:
    from xml.sax.saxutils import escape

    from reportlab.platypus import HRFlowable

    output: list[Any] = []
    for node in nodes:
        if node.tag == "p":
            if text_content(node.children).strip() == "\\pagebreak":
                output.append(PageBreak())
            elif text_content(node.children).strip():
                output.append(Paragraph(_inline_reportlab(node.children), styles["Body"]))
        elif node.tag.startswith("h") and node.tag[1:].isdigit():
            level = min(int(node.tag[1]), 3)
            output.append(Paragraph(_inline_reportlab(node.children), styles[f"Heading{level}"]))
        elif node.tag == "pre":
            output.append(Preformatted(text_content(node.children), styles["CodeBlock"]))
        elif node.tag in {"ul", "ol"}:
            ordered = node.tag == "ol"
            item_number = int(node.attrs.get("start", "1"))
            for item in node.children:
                if not isinstance(item, Node) or item.tag != "li":
                    continue
                marker = f"{item_number}." if ordered else "•"
                item_number += 1
                for child in item.children:
                    if isinstance(child, str):
                        if child.strip():
                            output.append(
                                Paragraph(
                                    escape(child.strip()),
                                    styles["Body"],
                                    bulletText=marker,
                                )
                            )
                        continue
                    if not isinstance(child, Node):
                        continue
                    if child.tag == "p":
                        output.append(
                            Paragraph(
                                _inline_reportlab(child.children),
                                styles["Body"],
                                bulletText=marker,
                            )
                        )
                    elif child.tag in {"ul", "ol"}:
                        output.extend(_pdf_blocks([child], styles, depth=depth + 1))
                    else:
                        output.extend(_pdf_blocks([child], styles, depth=depth + 1))
        elif node.tag == "table":
            rows = []
            for row in _table_rows(node):
                rows.append(
                    [Paragraph(_inline_reportlab(cell.children), styles["Body"]) for cell in row]
                )
            if rows:
                output.append(
                    Table(
                        rows,
                        repeatRows=1,
                        hAlign="LEFT",
                        style=TableStyle(
                            [
                                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf0f2")),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                                ("TOPPADDING", (0, 0), (-1, -1), 4),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                            ]
                        ),
                    )
                )
                output.append(Spacer(1, 7))
        elif node.tag == "blockquote":
            output.extend(
                _pdf_blocks(
                    [child for child in node.children if isinstance(child, Node)],
                    styles,
                    depth=depth + 1,
                )
            )
        elif node.tag == "hr":
            output.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#9aa0a6")))
            output.append(Spacer(1, 5))
        else:
            output.extend(
                _pdf_blocks(
                    [child for child in node.children if isinstance(child, Node)],
                    styles,
                    depth=depth,
                )
            )
    return output


def _table_rows(node: Node) -> list[list[Node]]:
    rows: list[list[Node]] = []

    def visit(current: Node) -> None:
        for child in current.children:
            if not isinstance(child, Node):
                continue
            if child.tag == "tr":
                rows.append(
                    [
                        cell
                        for cell in child.children
                        if isinstance(cell, Node) and cell.tag in {"th", "td"}
                    ]
                )
            else:
                visit(child)

    visit(node)
    return rows


def _append_docx_inline(
    paragraph: Any,
    content: list[Content],
    *,
    bold: bool = False,
    italic: bool = False,
    code: bool = False,
) -> None:
    for item in content:
        if isinstance(item, str):
            if item:
                run = paragraph.add_run(item)
                run.bold, run.italic = bold, italic
                if code:
                    run.font.name = "Consolas"
            continue
        _append_docx_inline(
            paragraph,
            item.children,
            bold=bold or item.tag in {"strong", "b"},
            italic=italic or item.tag in {"em", "i"},
            code=code or item.tag == "code",
        )
        if item.tag == "br":
            paragraph.add_run().add_break()


def _docx_blocks(nodes: list[Node], document: Any, *, depth: int = 0) -> None:
    from docx.shared import Inches, Pt

    for node in nodes:
        if node.tag == "p":
            if text_content(node.children).strip() == "\\pagebreak":
                document.add_page_break()
            elif text_content(node.children).strip():
                paragraph = document.add_paragraph()
                _append_docx_inline(paragraph, node.children)
        elif node.tag.startswith("h") and node.tag[1:].isdigit():
            paragraph = document.add_heading(level=min(int(node.tag[1]), 9))
            _append_docx_inline(paragraph, node.children)
        elif node.tag == "pre":
            paragraph = document.add_paragraph(style="No Spacing")
            for line_index, line in enumerate(text_content(node.children).splitlines()):
                if line_index:
                    paragraph.add_run().add_break()
                run = paragraph.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
        elif node.tag in {"ul", "ol"}:
            style = "List Bullet" if node.tag == "ul" else "List Number"
            for item in node.children:
                if not isinstance(item, Node) or item.tag != "li":
                    continue
                paragraph = document.add_paragraph(style=style)
                paragraph.paragraph_format.left_indent = Inches(0.25 * depth)
                for child in item.children:
                    if isinstance(child, Node) and child.tag in {"ul", "ol"}:
                        continue
                    if isinstance(child, Node) and child.tag == "p":
                        _append_docx_inline(paragraph, child.children)
                    elif isinstance(child, str):
                        _append_docx_inline(paragraph, [child])
                nested = [
                    child
                    for child in item.children
                    if isinstance(child, Node) and child.tag in {"ul", "ol"}
                ]
                _docx_blocks(nested, document, depth=depth + 1)
        elif node.tag == "table":
            rows = _table_rows(node)
            if rows:
                table = document.add_table(rows=0, cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                for source_row in rows:
                    cells = table.add_row().cells
                    for index, source_cell in enumerate(source_row):
                        _append_docx_inline(cells[index].paragraphs[0], source_cell.children)
        elif node.tag == "blockquote":
            for child in node.children:
                if isinstance(child, Node):
                    _docx_blocks([child], document, depth=depth)
                    if child.tag == "p":
                        document.paragraphs[-1].style = "Quote"
        elif node.tag == "hr":
            document.add_paragraph("────────────────────────────────")
        else:
            _docx_blocks(
                [child for child in node.children if isinstance(child, Node)], document, depth=depth
            )


def _document_sections(
    template: Template, solution: dict[str, Any]
) -> list[tuple[str | None, str]]:
    sections: list[tuple[str | None, str]] = []
    items = solution.get("question_answers", [])
    if template == Template.QUESTION_ANSWER:
        if items:
            for index, item in enumerate(items, 1):
                label = str(item.get("question_id") or f"Questão {index}")
                sections.append((f"{index}. {label}", str(item.get("answer", ""))))
        elif solution.get("answer"):
            sections.append((None, str(solution["answer"])))
        return sections
    if solution.get("understanding"):
        sections.append(("Entendimento", str(solution["understanding"])))
    if solution.get("answer"):
        sections.append(("Resposta", str(solution["answer"])))
    for index, item in enumerate(items, 1):
        label = str(item.get("question_id") or f"Questão {index}")
        sections.append((f"{index}. {label}", str(item.get("answer", ""))))
    if template == Template.CODE_ASSIGNMENT:
        for spec in solution.get("artifacts", []):
            if not isinstance(spec, dict):
                continue
            title = str(spec.get("title") or "Código")
            if spec.get("kind") == "CODE" or "code" in title.lower():
                source = str(spec.get("specification") or spec.get("content") or "")
                fenced = re.fullmatch(
                    r"(`{3,}|~{3,})[A-Za-z0-9_+.-]*[ \t]*\r?\n([\s\S]*?)\r?\n\1[ \t]*",
                    source,
                )
                if fenced:
                    source = fenced.group(2)
                longest_ticks = max((len(run) for run in re.findall(r"`+", source)), default=2)
                fence = "`" * max(3, longest_ticks + 1)
                code = source + ("" if source.endswith("\n") else "\n")
                sections.append((title, f"{fence}text\n{code}{fence}"))
    return sections


class DocumentBuilder:
    def __init__(self, engine: Engine, settings: Settings) -> None:
        self.engine, self.settings = engine, settings

    def choose_template(self, solution: dict[str, Any], override: str | None = None) -> Template:
        if override:
            try:
                return Template(override)
            except ValueError as exc:
                raise AppError(
                    "Template inválido; use question-answer, academic-report ou code-assignment."
                ) from exc
        types = set(solution.get("assignment_types", []))
        if "PROGRAMMING" in types:
            return Template.CODE_ASSIGNMENT
        if types & {"LONG_FORM", "RESEARCH"}:
            return Template.ACADEMIC_REPORT
        return Template.QUESTION_ANSWER

    @staticmethod
    def choose_format(override: str | OutputFormat | None) -> OutputFormat:
        try:
            if override is None:
                return OutputFormat.PDF
            if not isinstance(override, (str, OutputFormat)):
                raise ValueError
            return OutputFormat(str(override).lower())
        except (TypeError, ValueError) as exc:
            raise AppError("Formato inválido; use pdf, docx, txt ou md.") from exc

    def _assignment(self, assignment_id: int) -> tuple[AssignmentRecord, dict[str, Any]]:
        with Session(self.engine) as s:
            row = s.get(AssignmentRecord, assignment_id)
            if row is None or not row.present:
                raise AppError("Atividade local não encontrada.")
            course = s.get(
                __import__("app.persistence.models", fromlist=["CourseRecord"]).CourseRecord,
                row.course_id,
            )
            return row, (course.snapshot if course else {})

    def build(
        self,
        assignment_id: int,
        solution_version: int | None = None,
        override: str | None = None,
        output_format: str | OutputFormat | None = None,
    ) -> GeneratedArtifact:
        selected_format = self.choose_format(output_format)
        assignment, course = self._assignment(assignment_id)
        with Session(self.engine) as s:
            q = select(SolutionRecord).where(
                SolutionRecord.assignment_id == assignment_id,
                SolutionRecord.status.in_(("READY", "NEEDS_REVIEW", "APPROVED")),
            )
            if solution_version is not None:
                q = q.where(SolutionRecord.version == solution_version)
            solution_row = s.scalars(q.order_by(SolutionRecord.version.desc())).first()
            if solution_row is None:
                raise AppError("Nenhuma solução READY/NEEDS_REVIEW disponível para gerar PDF.")
            solution = solution_row.response or {
                "answer": solution_row.answer or "",
                "assignment_types": ["UNKNOWN"],
                "question_answers": [],
            }
            template = self.choose_template(solution, override)
            current = select(func.max(GeneratedArtifact.version)).where(
                GeneratedArtifact.assignment_id == assignment_id
            )
            version = (s.scalar(current) or 0) + 1
            title = _safe(assignment.snapshot.get("title", "Atividade"), "Atividade")
            subject = _safe(course.get("name", self.settings.course_name), "Disciplina")
            folder = (
                self.settings.generated_root.resolve()
                / _safe(str(course.get("id", assignment.course_id)))
                / str(assignment_id)
            )
            folder.mkdir(parents=True, exist_ok=True)
            base = f"{title} - {subject}"
            extension = FORMAT_EXTENSIONS[selected_format]
            filename = _safe(base, "atividade") + (
                extension if version == 1 else f" - v{version}{extension}"
            )
            path = folder / filename
            while path.exists():
                version += 1
                path = folder / (_safe(base, "atividade") + f" - v{version}{extension}")
            artifact = GeneratedArtifact(
                solution_id=solution_row.id,
                assignment_id=assignment_id,
                artifact_type=FORMAT_ARTIFACT_TYPES[selected_format],
                template=template.name,
                version=version,
                local_path=str(path),
                content_hash="0" * 64,
                status="GENERATING",
                created_at=datetime.now(UTC),
                updated_at=datetime.now(UTC),
            )
            s.add(artifact)
            s.commit()
            s.refresh(artifact)
        try:
            self._render(path, template, solution, assignment.snapshot, course, selected_format)
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            with Session(self.engine) as s:
                saved = s.get(GeneratedArtifact, artifact.id)
                assert saved
                saved.content_hash, saved.status, saved.updated_at = (
                    digest,
                    "UPLOAD_PENDING",
                    datetime.now(UTC),
                )
                s.commit()
                s.refresh(saved)
                return saved
        except Exception as exc:
            with Session(self.engine) as s:
                saved = s.get(GeneratedArtifact, artifact.id)
                assert saved
                saved.status, saved.error, saved.updated_at = (
                    "FAILED",
                    f"Falha ao renderizar {selected_format.value.upper()}",
                    datetime.now(UTC),
                )
                s.commit()
            raise AppError(f"Falha segura ao gerar {selected_format.value.upper()}.") from exc

    def build_all(
        self,
        assignment_id: int,
        solution_version: int | None = None,
        override: str | None = None,
        output_format: str | OutputFormat | None = None,
    ) -> list[GeneratedArtifact]:
        """Create the selected review document plus task-requested concrete files."""
        selected_format = self.choose_format(output_format)
        main_artifact = self.build(assignment_id, solution_version, override, selected_format)
        with Session(self.engine) as s:
            solution_row = s.get(SolutionRecord, main_artifact.solution_id)
            solution = solution_row.response if solution_row else None
        output = [main_artifact]
        if not solution:
            return output
        for spec in solution.get("artifacts", []):
            if not isinstance(spec, dict):
                continue
            name = _deliverable_name(spec.get("filename") or spec.get("title"))
            content = _deliverable_content(spec.get("content") or spec.get("specification"))
            if name and content:
                output.append(
                    self._write_deliverable(assignment_id, main_artifact.solution_id, name, content)
                )
        return output

    def _write_deliverable(
        self, assignment_id: int, solution_id: int, name: str, content: str
    ) -> GeneratedArtifact:
        data = content.encode("utf-8")
        if len(data) > 2_000_000:
            raise AppError(f"Artefato {name} excede o limite de 2 MB.")
        assignment, course = self._assignment(assignment_id)
        folder = (
            self.settings.generated_root.resolve()
            / _safe(str(course.get("id", assignment.course_id)))
            / str(assignment_id)
        )
        folder.mkdir(parents=True, exist_ok=True)
        with Session(self.engine) as s:
            version = (
                s.scalar(
                    select(func.max(GeneratedArtifact.version)).where(
                        GeneratedArtifact.assignment_id == assignment_id
                    )
                )
                or 0
            ) + 1
            path = folder / name
            if path.exists():
                path = folder / f"{Path(name).stem} - v{version}{Path(name).suffix}"
            now = datetime.now(UTC)
            artifact = GeneratedArtifact(
                solution_id=solution_id,
                assignment_id=assignment_id,
                artifact_type=TEXT_DELIVERABLES[Path(name).suffix.lower()],
                template="DELIVERABLE",
                version=version,
                local_path=str(path),
                content_hash=hashlib.sha256(data).hexdigest(),
                status="GENERATING",
                created_at=now,
                updated_at=now,
            )
            s.add(artifact)
            s.commit()
            s.refresh(artifact)
        try:
            path.write_bytes(data)
            with Session(self.engine) as s:
                saved = s.get(GeneratedArtifact, artifact.id)
                assert saved
                saved.status = "UPLOAD_PENDING"
                saved.updated_at = datetime.now(UTC)
                s.commit()
                s.refresh(saved)
                return saved
        except OSError as exc:
            with Session(self.engine) as s:
                saved = s.get(GeneratedArtifact, artifact.id)
                assert saved
                saved.status = "FAILED"
                saved.error = "Falha ao gravar artefato solicitado"
                saved.updated_at = datetime.now(UTC)
                s.commit()
            raise AppError(f"Falha segura ao gerar {name}.") from exc

    def _render(
        self,
        path: Path,
        template: Template,
        solution: dict[str, Any],
        assignment: dict[str, Any],
        course: dict[str, Any],
        output_format: OutputFormat = OutputFormat.PDF,
    ) -> None:
        if output_format == OutputFormat.PDF:
            self._render_pdf(path, template, solution, assignment, course)
        elif output_format == OutputFormat.DOCX:
            self._render_docx(path, template, solution, assignment, course)
        elif output_format == OutputFormat.TXT:
            path.write_text(
                self._plain_document(template, solution, assignment, course), encoding="utf-8"
            )
        elif output_format == OutputFormat.MD:
            path.write_text(
                self._markdown_document(template, solution, assignment, course), encoding="utf-8"
            )

    def _header(self, assignment: dict[str, Any], course: dict[str, Any]) -> tuple[str, list[str]]:
        title = str(assignment.get("title") or "ATIVIDADE")
        metadata = [f"Data: {datetime.now(self.settings.zone):%d/%m/%Y}"]
        if self.settings.student_name:
            metadata.append(f"Aluno: {self.settings.student_name}")
        metadata.append(f"Matéria: {course.get('name', self.settings.course_name)}")
        return title, metadata

    def _plain_document(
        self,
        template: Template,
        solution: dict[str, Any],
        assignment: dict[str, Any],
        course: dict[str, Any],
    ) -> str:
        title, metadata = self._header(assignment, course)
        parts = [title, *metadata]
        for heading, markdown in _document_sections(template, solution):
            if heading:
                parts.extend(["", heading])
            parts.extend(["", clean_text(markdown).rstrip()])
        return "\n".join(parts).strip() + "\n"

    def _markdown_document(
        self,
        template: Template,
        solution: dict[str, Any],
        assignment: dict[str, Any],
        course: dict[str, Any],
    ) -> str:
        title, metadata = self._header(assignment, course)
        parts = [f"# {title}", "", *(f"_{line}_  " for line in metadata), ""]
        for heading, markdown in _document_sections(template, solution):
            if heading:
                parts.extend([f"## {heading}", ""])
            parts.extend([markdown.rstrip(), ""])
        return "\n".join(parts).rstrip() + "\n"

    def _render_pdf(
        self,
        path: Path,
        template: Template,
        solution: dict[str, Any],
        assignment: dict[str, Any],
        course: dict[str, Any],
    ) -> None:
        styles = getSampleStyleSheet()
        styles.add(
            ParagraphStyle(
                name="AcademicTitle",
                parent=styles["Title"],
                alignment=TA_CENTER,
                fontSize=16,
                leading=20,
                spaceAfter=14,
            )
        )
        styles.add(
            ParagraphStyle(
                name="Section",
                parent=styles["Heading2"],
                fontSize=13,
                leading=16,
                spaceBefore=10,
                spaceAfter=6,
            )
        )
        styles.add(
            ParagraphStyle(
                name="Body", parent=styles["BodyText"], fontSize=11, leading=16, spaceAfter=7
            )
        )
        styles.add(
            ParagraphStyle(
                name="Meta", parent=styles["BodyText"], fontSize=10, leading=14, alignment=TA_CENTER
            )
        )
        styles.add(
            ParagraphStyle(
                name="Question",
                parent=styles["BodyText"],
                fontSize=11,
                leading=15,
                spaceBefore=8,
                spaceAfter=3,
            )
        )
        styles.add(
            ParagraphStyle(
                name="CodeBlock",
                fontName="Courier",
                fontSize=8.5,
                leading=10.5,
                leftIndent=10,
                rightIndent=10,
                spaceBefore=6,
                spaceAfter=8,
            )
        )
        doc = SimpleDocTemplate(
            str(path),
            pagesize=A4,
            rightMargin=2 * cm,
            leftMargin=2 * cm,
            topMargin=1.8 * cm,
            bottomMargin=1.8 * cm,
            title=str(assignment.get("title", "Atividade")),
            author=self.settings.student_name or "Classroom Agent",
        )
        story: list[Any] = []
        title, metadata = self._header(assignment, course)
        title = _esc(title)
        if template == Template.CODE_ASSIGNMENT:
            for value in (
                self.settings.institution,
                self.settings.unit,
                self.settings.course_name,
                f"Discente: {self.settings.student_name}" if self.settings.student_name else "",
            ):
                if value:
                    story.append(Paragraph(_esc(value), styles["Meta"]))
            story += [Spacer(1, 14), Paragraph(title, styles["AcademicTitle"])]
        else:
            story.append(Paragraph(title, styles["AcademicTitle"]))
            for value in metadata:
                story.append(Paragraph(_esc(value), styles["Meta"]))
        sections = _document_sections(template, solution)
        for heading, markdown in sections:
            if heading:
                story.append(Paragraph(_esc(heading), styles["Section"]))
            story.extend(_pdf_blocks(parse_markdown(markdown), styles))
        doc.build(story, onFirstPage=self._footer, onLaterPages=self._footer)

    def _render_docx(
        self,
        path: Path,
        template: Template,
        solution: dict[str, Any],
        assignment: dict[str, Any],
        course: dict[str, Any],
    ) -> None:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt

        document = Document()
        title, metadata = self._header(assignment, course)
        document.add_heading(title, level=0)
        for value in metadata:
            paragraph = document.add_paragraph(value)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for heading, markdown in _document_sections(template, solution):
            if heading:
                document.add_heading(heading, level=1)
            nodes = parse_markdown(markdown)
            _docx_blocks(nodes, document)
        for section in document.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
        document.styles["Normal"].font.name = "Calibri"
        document.styles["Normal"].font.size = Pt(11)
        document.save(str(path))

    @staticmethod
    def _footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 9)
        canvas.drawCentredString(A4[0] / 2, 1 * cm, f"Página {doc.page}")
        canvas.restoreState()
